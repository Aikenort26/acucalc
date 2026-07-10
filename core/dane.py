"""Proyecciones de población municipal DANE (dataset embebido data/dane.csv.gz).

El DANE publica proyecciones por municipio y área geográfica (Cabecera Municipal,
Centros Poblados y Rural Disperso, Total) hasta 2042 y las actualiza
periódicamente; `update_from_file` permite reemplazar el dataset embebido.
"""
import datetime as dt
import gzip
import json
from pathlib import Path

import pandas as pd

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
CSV = DATA_DIR / "dane.csv.gz"
META = DATA_DIR / "dane_meta.json"

REQUIRED = ["DP", "DPNOM", "DPMP", "MPIO", "AÑO", "ÁREA GEOGRÁFICA", "Población"]

_cache: dict = {}


def load() -> pd.DataFrame:
    """DataFrame DPNOM/MPIO/ANO/AREA/POBLACION, cacheado en memoria."""
    if "df" not in _cache:
        _cache["df"] = pd.read_csv(CSV, dtype={"DP": str, "DPMP": str})
    return _cache["df"]


def meta() -> dict:
    return json.loads(META.read_text(encoding="utf-8"))


def departamentos() -> list[str]:
    return sorted(load()["DPNOM"].unique())


def municipios(dpto: str) -> list[str]:
    df = load()
    return sorted(df.loc[df["DPNOM"] == dpto, "MPIO"].unique())


def areas() -> list[str]:
    return sorted(load()["AREA"].unique())


def series(dpto: str, mpio: str, area: str) -> list[tuple[int, int]]:
    """Serie (año, población) ordenada para un municipio y área."""
    df = load()
    sel = df[(df["DPNOM"] == dpto) & (df["MPIO"] == mpio) & (df["AREA"] == area)]
    return [(int(r.ANO), int(r.POBLACION)) for r in
            sel.sort_values("ANO").itertuples()]


def _find_header_row(rows) -> int | None:
    for i, row in enumerate(rows):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if cells[: len(REQUIRED)] == REQUIRED:
            return i
    return None


def update_from_file(path: str | Path) -> dict:
    """Reemplaza el dataset embebido con un archivo DANE nuevo (xlsx o csv).

    Autodetecta la fila de encabezado DP|DPNOM|DPMP|MPIO|AÑO|ÁREA GEOGRÁFICA|Población.
    Lanza ValueError si el formato no coincide.
    """
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        for name in wb.sheetnames:
            if name.strip().upper() == "DANE":
                ws = wb[name]
                break
        all_rows = list(ws.iter_rows(values_only=True))
        wb.close()
    else:
        raw = pd.read_csv(path, header=None, dtype=str)
        all_rows = [tuple(r) for r in raw.itertuples(index=False)]

    h = _find_header_row(all_rows)
    if h is None:
        raise ValueError(f"El archivo no contiene el encabezado DANE esperado: {REQUIRED}")

    out_rows, years, mpios = [], set(), set()
    for row in all_rows[h + 1:]:
        vals = (tuple(row) + (None,) * 7)[:7]
        dp, dpnom, dpmp, mpio, ano, area, pob = vals
        if dpnom is None or ano is None or pob is None or str(pob).strip() == "":
            continue
        try:
            ano_i, pob_i = int(float(ano)), int(float(pob))
        except (TypeError, ValueError):
            continue
        out_rows.append((str(dp), str(dpnom).strip(), str(dpmp), str(mpio).strip(),
                         ano_i, str(area).strip(), pob_i))
        years.add(ano_i)
        mpios.add((str(dpnom).strip(), str(mpio).strip()))
    if not out_rows:
        raise ValueError("El archivo DANE no contiene filas de datos válidas")

    with gzip.open(CSV, "wt", encoding="utf-8", newline="") as f:
        f.write("DP,DPNOM,DPMP,MPIO,ANO,AREA,POBLACION\n")
        for r in out_rows:
            f.write(",".join(f'"{v}"' if "," in str(v) else str(v) for v in r) + "\n")

    new_meta = {
        "fuente": path.name,
        "fecha_extraccion": dt.date.today().isoformat(),
        "anos": [min(years), max(years)],
        "municipios": len(mpios),
        "filas": len(out_rows),
    }
    META.write_text(json.dumps(new_meta, ensure_ascii=False, indent=1), encoding="utf-8")
    _cache.clear()
    return new_meta
